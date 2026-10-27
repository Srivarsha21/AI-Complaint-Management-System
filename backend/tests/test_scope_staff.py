"""Department staff get their department's complaints in reports (merge of modules 4 and 7)."""

from io import BytesIO

from openpyxl import load_workbook

from app.models import UserRole
from tests.conftest import auth_header, login, make_category, make_user

COMPLAINTS = "/api/v1/complaints"
VALID = {
    "subject": "Blocked drain near bus stop",
    "description": "The drain outside the bus stop has been blocked for a week and water is overflowing.",
    "location_text": "Ameerpet Main Road, Hyderabad",
}


def test_staff_report_contains_only_their_department(client, db):
    drainage = make_category(db, name="Drainage", slug="drainage", department="Sewerage Dept")
    roads = make_category(db, name="Roads", slug="roads", department="Roads Dept")
    citizen = make_user(db, email="c@example.com")
    staff = make_user(db, email="s@example.com", role=UserRole.STAFF)
    staff.department_id = drainage.department_id
    db.commit()

    tokens = login(client, "c@example.com")
    for category in (drainage, roads):
        response = client.post(COMPLAINTS, data={**VALID, "category_id": str(category.id)}, headers=auth_header(tokens))
        assert response.status_code == 201, response.text

    report = client.get("/api/v1/reports/complaints.xlsx", headers=auth_header(login(client, "s@example.com")))
    assert report.status_code == 200
    sheet = load_workbook(BytesIO(report.content))["Complaints"]
    rows = [row for row in sheet.iter_rows(min_row=2, values_only=True) if any(row)]
    # Only the drainage complaint; the roads complaint belongs to another department.
    assert len(rows) == 1
    assert citizen.id is not None
