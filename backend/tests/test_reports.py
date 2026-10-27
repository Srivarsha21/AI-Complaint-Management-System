import base64
import re
import zlib
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import AuditLog, Complaint
from app.services import report_service
from tests import analytics_fixtures
from tests.conftest import auth_header, login

PDF = "/api/v1/reports/complaints.pdf"
XLSX = "/api/v1/reports/complaints.xlsx"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ALL_REFS = {f"CMP-2026-{n:06d}" for n in range(1, 7)}


@pytest.fixture
def data(db):
    return analytics_fixtures.build(db)


def _download(client, email, url, **params):
    response = client.get(url, params=params, headers=auth_header(login(client, email)))
    assert response.status_code == 200, response.text
    return response


def _pdf_text(content: bytes) -> str:
    """Decompress every content stream so drawn text can be searched."""
    chunks = []
    for match in re.finditer(rb"/Filter\s*(\[[^\]]*\]|/\w+).*?stream\r?\n(.*?)endstream", content, re.S):
        filters, raw = match.group(1), match.group(2).strip()
        try:
            if b"ASCII85Decode" in filters:
                raw = base64.a85decode(raw[:-2] if raw.endswith(b"~>") else raw)
            if b"FlateDecode" in filters:
                raw = zlib.decompress(raw)
        except (ValueError, zlib.error):
            continue
        chunks.append(raw)
    return b"\n".join(chunks).decode("latin-1")


def _workbook(response):
    return load_workbook(BytesIO(response.content))


def _rows(ws):
    header = [c.value for c in ws[1]]
    return header, [dict(zip(header, (c.value for c in row))) for row in ws.iter_rows(min_row=2)]


def _all_text(wb) -> str:
    return "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)


# --- PDF -------------------------------------------------------------------------

def test_admin_pdf_contains_all_references_and_is_audited(client, data, db):
    response = _download(client, "admin@example.com", PDF)
    assert response.headers["content-type"] == "application/pdf"
    assert re.fullmatch(r'attachment; filename="complaints-\d{4}-\d{2}-\d{2}\.pdf"',
                        response.headers["content-disposition"])
    assert response.content.startswith(b"%PDF")
    text = _pdf_text(response.content)
    for ref in ALL_REFS:
        assert ref in text
    assert "Alice Reddy" in text  # complainant column for admins
    assert "alice@example.com" not in text
    assert "Page 1 of" in text

    entry = db.query(AuditLog).filter_by(action="report.exported").one()
    assert entry.actor_id == data.admin.id
    assert entry.details["format"] == "pdf"
    assert entry.details["rows"] == 6


def test_user_pdf_contains_only_own_complaints(client, data, db):
    response = _download(client, "alice@example.com", PDF)
    text = _pdf_text(response.content)
    assert {ref for ref in ALL_REFS if ref in text} == {data.ref(1), data.ref(2), data.ref(6)}
    assert "Complainant" not in text
    assert "Your complaints only" in text
    # Only administrator exports are audited.
    assert db.query(AuditLog).filter_by(action="report.exported").count() == 0


def test_pdf_with_no_matching_complaints(client, data):
    response = _download(client, "admin@example.com", PDF, date_from="2020-01-01", date_to="2020-01-02")
    assert response.content.startswith(b"%PDF")
    assert "No complaints match these filters." in _pdf_text(response.content)


# --- Excel -----------------------------------------------------------------------

def test_admin_xlsx_matches_database(client, data, db):
    response = _download(client, "admin@example.com", XLSX)
    assert response.headers["content-type"] == XLSX_TYPE
    assert re.fullmatch(r'attachment; filename="complaints-\d{4}-\d{2}-\d{2}\.xlsx"',
                        response.headers["content-disposition"])

    wb = _workbook(response)
    assert wb.sheetnames == ["Complaints", "Summary"]
    ws = wb["Complaints"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:J7"
    assert all(c.font.bold for c in ws[1])

    header, rows = _rows(ws)
    assert header == ["Reference", "Subject", "Category", "Department", "Priority", "Status", "Complainant",
                      "Created (IST)", "Resolved (IST)", "Closed (IST)"]
    by_ref = {r["Reference"]: r for r in rows}
    assert set(by_ref) == ALL_REFS

    for complaint in db.query(Complaint).all():
        row = by_ref[complaint.reference_number]
        assert row["Subject"] == complaint.subject
        assert row["Category"] == complaint.category.name
        assert row["Complainant"] == complaint.user.full_name
        assert isinstance(row["Created (IST)"], datetime)
        expected = complaint.created_at.astimezone(analytics_fixtures.IST).replace(tzinfo=None)
        assert row["Created (IST)"] == expected

    assert by_ref[data.ref(5)]["Department"] == "Unassigned"
    assert by_ref[data.ref(2)]["Priority"] == "Not assessed"
    assert by_ref[data.ref(4)]["Closed (IST)"] == datetime(2026, 10, 5, 1, 0)  # 3 Oct 09:00 + 40 h

    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=3) if row[0].value}
    assert summary["Matching complaints"] == 6
    assert summary["Resolved"] == 2
    assert summary["Drainage"] == 3 and summary["Potholes"] == 3

    text = _all_text(wb)
    assert "@example.com" not in text

    entry = db.query(AuditLog).filter_by(action="report.exported").one()
    assert entry.details["format"] == "xlsx"


def test_user_xlsx_has_only_own_complaints_and_no_complainant_column(client, data, db):
    wb = _workbook(_download(client, "bob@example.com", XLSX))
    header, rows = _rows(wb["Complaints"])
    assert "Complainant" not in header
    assert {r["Reference"] for r in rows} == {data.ref(3), data.ref(4), data.ref(5)}
    text = _all_text(wb)
    assert "Bob Naidu" not in text and "Alice Reddy" not in text and "@example.com" not in text
    assert db.query(AuditLog).filter_by(action="report.exported").count() == 0


def test_report_filters_and_audit_details(client, data, db):
    params = {"status": "resolved", "date_from": "2026-09-01", "date_to": "2026-10-31",
              "category_id": data.drainage.id, "search": "number 1"}
    wb = _workbook(_download(client, "admin@example.com", XLSX, **params))
    _, rows = _rows(wb["Complaints"])
    assert [r["Reference"] for r in rows] == [data.ref(1)]

    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=3) if row[0].value}
    assert summary["Filter: Category"] == "Drainage"
    assert summary["Filter: Status"] == "Resolved"
    assert summary["Filter: Search"] == "number 1"

    details = db.query(AuditLog).filter_by(action="report.exported").one().details
    assert details["filters"] == {
        "date_from": "2026-09-01", "date_to": "2026-10-31", "category_id": data.drainage.id,
        "department_id": None, "status": "resolved", "search": "number 1",
    }


def test_users_cannot_widen_scope_with_filters(client, data):
    wb = _workbook(_download(client, "alice@example.com", XLSX, department_id=data.drainage.department_id))
    _, rows = _rows(wb["Complaints"])
    assert [r["Reference"] for r in rows] == [data.ref(1)]


def test_rows_are_capped_with_a_note(client, data, monkeypatch):
    monkeypatch.setattr(report_service, "MAX_REPORT_ROWS", 2)
    wb = _workbook(_download(client, "admin@example.com", XLSX))
    _, rows = _rows(wb["Complaints"])
    assert len(rows) == 2
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=3) if row[0].value}
    assert summary["Matching complaints"] == 6
    assert summary["Rows in this report"] == 2
    assert "Showing the newest 2 of 6" in summary["Note"]

    pdf = _download(client, "admin@example.com", PDF)
    assert "Showing the newest 2 of 6" in _pdf_text(pdf.content)


def test_text_starting_with_equals_is_not_a_formula(client, data, db):
    complaint = data.complaints[data.ref(1)]
    complaint.subject = '=HYPERLINK("http://evil.example","click")'
    db.commit()
    ws = _workbook(_download(client, "admin@example.com", XLSX))["Complaints"]
    cell = next(row[1] for row in ws.iter_rows(min_row=2) if row[0].value == data.ref(1))
    assert cell.data_type == "s"
    assert cell.value == complaint.subject


def test_pdf_escapes_markup_in_subjects(client, data, db):
    data.complaints[data.ref(2)].subject = "Broken <b>pipe</b> & leak"
    db.commit()
    response = _download(client, "admin@example.com", PDF)
    assert response.content.startswith(b"%PDF")


def test_reports_require_authentication(client, data):
    assert client.get(PDF).status_code == 401
    assert client.get(XLSX).status_code == 401
