"""Complaint report as an Excel workbook (OpenPyXL).

Sheets:
  * "Complaints": bold, frozen header row with an autofilter; dates are real
    Excel date-time cells in IST (Excel has no time zones, so the header says IST).
  * "Summary": generation details, applied filters and counts by status/category.
"""

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.common import PRIORITY_LABELS, STATUS_LABELS, ReportData, to_ist

DATE_FORMAT = "dd-mmm-yyyy hh:mm"
BOLD = Font(bold=True)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="312E81")


def _naive_ist(value):
    local = to_ist(value)
    return local.replace(tzinfo=None) if local else None


def _set(ws: Worksheet, row: int, column: int, value: Any):
    cell = ws.cell(row=row, column=column, value=value)
    # Text that starts with "=" must never become a formula (spreadsheet injection).
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def _complaints_sheet(ws: Worksheet, report: ReportData) -> None:
    columns: list[tuple[str, int]] = [
        ("Reference", 18),
        ("Subject", 48),
        ("Category", 26),
        ("Department", 28),
        ("Priority", 13),
        ("Status", 13),
    ]
    if report.include_complainant:
        columns.append(("Complainant", 24))
    columns += [("Created (IST)", 18), ("Resolved (IST)", 18), ("Closed (IST)", 18)]

    for index, (name, width) in enumerate(columns, start=1):
        cell = _set(ws, 1, index, name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(index)].width = width

    for r, row in enumerate(report.rows, start=2):
        values: list[Any] = [
            row.reference_number,
            row.subject,
            row.category,
            row.department or "Unassigned",
            PRIORITY_LABELS.get(row.priority, row.priority),
            STATUS_LABELS.get(row.status, row.status),
        ]
        if report.include_complainant:
            values.append(row.complainant)
        values += [_naive_ist(row.created_at), _naive_ist(row.resolved_at), _naive_ist(row.closed_at)]
        for c, value in enumerate(values, start=1):
            cell = _set(ws, r, c, value)
            if c > len(values) - 3:
                cell.number_format = DATE_FORMAT

    ws.freeze_panes = "A2"
    last_row = max(1, len(report.rows) + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"


def _summary_sheet(ws: Worksheet, report: ReportData) -> None:
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60

    title = _set(ws, 1, 1, report.title)
    title.font = Font(bold=True, size=14)
    rows: list[tuple[str, Any]] = [
        ("Generated (IST)", _naive_ist(report.generated_at)),
        ("Scope", report.scope_label),
        *[(f"Filter: {label}", value) for label, value in report.filters],
        ("Matching complaints", report.total_matching),
        ("Rows in this report", len(report.rows)),
    ]
    if report.truncation_note:
        rows.append(("Note", report.truncation_note))

    r = 3
    for label, value in rows:
        _set(ws, r, 1, label).font = BOLD
        cell = _set(ws, r, 2, value)
        cell.alignment = Alignment(horizontal="left")
        if label == "Generated (IST)":
            cell.number_format = DATE_FORMAT
        r += 1

    for heading, pairs in (("Status", report.status_counts), ("Category", report.category_counts)):
        r += 1
        _set(ws, r, 1, heading).font = BOLD
        _set(ws, r, 2, "Count").font = BOLD
        r += 1
        for label, count in pairs:
            _set(ws, r, 1, label)
            _set(ws, r, 2, count).alignment = Alignment(horizontal="left")
            r += 1


def build_xlsx(report: ReportData) -> bytes:
    workbook = Workbook()
    complaints = workbook.active
    complaints.title = "Complaints"
    _complaints_sheet(complaints, report)
    _summary_sheet(workbook.create_sheet("Summary"), report)
    workbook.properties.title = report.title
    workbook.properties.creator = "ResolveAI"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
